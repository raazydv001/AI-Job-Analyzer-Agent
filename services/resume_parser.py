import io
import zipfile


MAX_FILE_SIZE_MB = 5
MAX_FILE_SIZE_BYTES = MAX_FILE_SIZE_MB * 1024 * 1024


def validate_resume_file(uploaded_file):
    """
    Validate the uploaded resume before attempting extraction.

    Returns:
        (True, None) when valid
        (False, error_message) when invalid
    """

    if uploaded_file is None:
        return False, "No resume file was uploaded."

    filename = uploaded_file.name

    if not filename:
        return False, "The uploaded file has no filename."

    extension = filename.lower().split(".")[-1]

    # Check extension
    if extension not in {"pdf", "docx"}:
        return (
            False,
            "Unsupported file type. Please upload a PDF or DOCX file."
        )

    # Check size
    file_size = uploaded_file.size

    if file_size == 0:
        return False, "The uploaded file is empty."

    if file_size > MAX_FILE_SIZE_BYTES:
        return (
            False,
            f"File is too large. Maximum allowed size is "
            f"{MAX_FILE_SIZE_MB} MB."
        )

    # Read bytes for content validation
    file_bytes = uploaded_file.getvalue()

    if extension == "pdf":
        if not file_bytes.startswith(b"%PDF"):
            return (
                False,
                "The file extension is .pdf, but the file does not "
                "appear to be a valid PDF."
            )

    elif extension == "docx":
        if not _is_valid_docx(file_bytes):
            return (
                False,
                "The file extension is .docx, but the file does not "
                "appear to be a valid DOCX document."
            )

    return True, None


def _is_valid_docx(file_bytes):
    """
    Check whether the uploaded bytes represent a valid DOCX package.

    DOCX files are ZIP-based Office Open XML documents.
    """

    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as zip_file:

            required_files = {
                "[Content_Types].xml",
                "word/document.xml"
            }

            file_names = set(zip_file.namelist())

            return required_files.issubset(file_names)

    except zipfile.BadZipFile:
        return False


def extract_pdf_text(file_bytes):
    """
    Extract text from every page of a PDF.
    """

    import pymupdf

    try:
        with pymupdf.open(stream=file_bytes, filetype="pdf") as document:

            pages = []

            for page_number, page in enumerate(document, start=1):

                page_text = page.get_text("text").strip()

                if page_text:
                    pages.append(
                        f"--- Page {page_number} ---\n"
                        f"{page_text}"
                    )

            text = "\n\n".join(pages)

            if not text.strip():
                raise ValueError(
                    "No text could be extracted from the PDF. "
                    "The PDF may be scanned/image-based."
                )

            return text.strip()

    except Exception as e:
        raise ValueError(
            f"Could not read the PDF: {str(e)}"
        ) from e


def extract_docx_text(file_bytes):
    """
    Extract text from paragraphs and tables in a DOCX file.
    """

    from docx import Document

    try:
        document = Document(
            io.BytesIO(file_bytes)
        )

        parts = []

        # Extract normal paragraphs
        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:
                parts.append(text)

        # Extract tables
        for table in document.tables:

            for row in table.rows:

                row_text = []

                for cell in row.cells:

                    cell_text = cell.text.strip()

                    if cell_text:
                        row_text.append(cell_text)

                if row_text:
                    parts.append(" | ".join(row_text))

        text = "\n".join(parts)

        if not text.strip():
            raise ValueError(
                "No text could be extracted from the DOCX file."
            )

        return text.strip()

    except Exception as e:
        raise ValueError(
            f"Could not read the DOCX file: {str(e)}"
        ) from e


def extract_resume_text(uploaded_file):
    """
    Main function used by Streamlit.

    Validates the uploaded file and extracts its text.
    """

    is_valid, error_message = validate_resume_file(
        uploaded_file
    )

    if not is_valid:
        raise ValueError(error_message)

    file_bytes = uploaded_file.getvalue()

    extension = uploaded_file.name.lower().split(".")[-1]

    if extension == "pdf":
        return extract_pdf_text(file_bytes)

    if extension == "docx":
        return extract_docx_text(file_bytes)

    raise ValueError(
        "Unsupported resume file format."
    )